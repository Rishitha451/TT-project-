from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()

# Slide 1 - Project Overview
slide_layout = prs.slide_layouts[0] # Title slide
slide1 = prs.slides.add_slide(slide_layout)
title1 = slide1.shapes.title
subtitle1 = slide1.placeholders[1]
title1.text = "AI Finance Advisor"
subtitle1.text = "Financial Management & Investment Intelligence Platform\n\n" \
                 "• Provides a centralized financial management experience\n" \
                 "• Adds structured investment intelligence\n" \
                 "• Combines portfolio, market, risk, and investment analysis\n" \
                 "• Provides a foundation for future AI-powered financial services\n\n" \
                 "Financial Management + Investment Intelligence = AI Finance Advisor"

# Slide 2 - Development Journey
slide_layout = prs.slide_layouts[1] # Title and Content
slide2 = prs.slides.add_slide(slide_layout)
title2 = slide2.shapes.title
title2.text = "From Finance Advisor to Investment Intelligence"
content2 = slide2.shapes.placeholders[1].text_frame
content2.text = "AI Finance Advisor\n↓\nTradeFlow Reference\n↓\nInvestment Foundation\n↓\nPortfolio\n↓\nMarket Intelligence\n↓\nRisk & Sentiment\n↓\nInvestment Insights\n↓\nUnified Investment Experience"
p2 = content2.add_paragraph()
p2.text = "\n• TradeFlow was used as a UX/reference source\n• Useful concepts were migrated\n• TradeFlow's random simulation behavior was removed\n• Everything was rebuilt inside the Finance Advisor architecture"

# Slide 3 - Investment Intelligence
slide3 = prs.slides.add_slide(slide_layout)
title3 = slide3.shapes.title
title3.text = "Investment Intelligence Module"
content3 = slide3.shapes.placeholders[1].text_frame
content3.text = "Portfolio\n- Holdings, value, P&L, allocation\n\n" \
                "Market Intelligence\n- News, sentiment, trends\n\n" \
                "Risk & Sentiment\n- Risk metrics, exposure, sentiment\n\n" \
                "Investment Insights\n- Structured analysis, confidence, reasoning\n\n" \
                "Portfolio ↓ Market Intelligence ↓ Risk & Sentiment ↓ Investment Insights"

# Slide 4 - Architecture & Engineering Validation
slide4 = prs.slides.add_slide(slide_layout)
title4 = slide4.shapes.title
title4.text = "Architecture & Engineering Validation"
content4 = slide4.shapes.placeholders[1].text_frame
content4.text = "Finance Advisor UI\n↓\nFeature Hooks\n↓\nTanStack Query\n↓\nAPI Adapter\n↓\nDeterministic Mock Services\n↓\nFuture FastAPI / AI / Data Services\n\n" \
                "Validation:\n✓ TypeScript\n✓ Production Build\n✓ Route Validation\n✓ Responsive UI\n✓ Runtime Validation\n✓ Simulation Audit\n✓ Regression Testing\n\n" \
                "TradeFlow reference project remained untouched.\n" \
                "The current application uses deterministic mock services. Backend and real AI/data integrations are planned for a later phase."

# Slide 5 - Final Outcome
slide5 = prs.slides.add_slide(slide_layout)
title5 = slide5.shapes.title
title5.text = "Final Outcome"
content5 = slide5.shapes.placeholders[1].text_frame
content5.text = "Completed:\n" \
                "• Complete investment module\n" \
                "• Unified Investment Hub\n" \
                "• Portfolio, Market, Risk, Insights, Dashboard integration, Responsive frontend, Validated architecture\n\n" \
                "Current Status:\n" \
                "Frontend: COMPLETE & DEMO READY\n\n" \
                "Next Stage (Future Work):\n" \
                "FastAPI Backend\n↓\nReal Financial Data\n↓\nAI / LLM Analysis\n↓\nProduction Data Contracts\n↓\nReal-Time Financial Intelligence"

prs.save("AI_Finance_Advisor_Phases_1-6.pptx")
